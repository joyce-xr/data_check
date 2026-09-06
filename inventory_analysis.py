import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ATTACH_DIR = r'D:\粤之星对账\26-8\成都鑫联昇'
BEFORE_INV = ATTACH_DIR + r'\7月-账面库存.xlsx'
NOW_FLOW = ATTACH_DIR + r'\8月-出入库流水.xlsx'
AFTER_INV = ATTACH_DIR + r'\8月-账面库存.xlsx'

OUTPUT = r'D:\粤之星对账\26-8\成都鑫联昇\8月库存差异分析.xlsx'

GROUP_KEYS = ['商品编码', '库存类型', '仓库', '货主']
NEG_PREFIXES = ['XSDD', 'CGTH', 'PKD', 'JSTH', 'DBCK', 'WTO', 'QTCK']

# ========== Styles ==========
header_font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
header_fill = PatternFill(fill_type='solid', fgColor='2B4570')
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

data_font = Font(name='Arial', size=9)
zebra1 = PatternFill(fill_type='solid', fgColor='FFFFFF')
zebra2 = PatternFill(fill_type='solid', fgColor='F2F5FA')
diff_fill = PatternFill(fill_type='solid', fgColor='FFF2CC')
diff_font = Font(name='Arial', size=9, bold=True, color='CC0000')

thin_border = Border(
    left=Side(style='thin', color='D9DEE7'),
    right=Side(style='thin', color='D9DEE7'),
    top=Side(style='thin', color='D9DEE7'),
    bottom=Side(style='thin', color='D9DEE7')
)
FMT_INT = '#,##0;(#,##0);-'
FMT_DEC = '#,##0.00;(#,##0.00);-'

def write_sheet(ws, headers, data_rows, int_cols, dec_cols, diff_col_start):
    """Write data to a worksheet with formatting."""
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(1, col_idx, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    for row_idx, row_data in enumerate(data_rows, 2):
        has_diff = row_data[-1]
        values = row_data[:-1]
        zebra = zebra1 if (row_idx % 2 == 0) else zebra2
        for col_idx, v in enumerate(values, 1):
            cell = ws.cell(row_idx, col_idx, v)
            is_diff_cell = has_diff and col_idx >= diff_col_start
            cell.font = diff_font if is_diff_cell else data_font
            cell.fill = diff_fill if is_diff_cell else zebra
            cell.border = thin_border
            if col_idx <= 4:
                cell.alignment = Alignment(horizontal='left', vertical='center')
            elif col_idx in int_cols:
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = FMT_INT
            elif col_idx in dec_cols:
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = FMT_DEC

    n_rows = len(data_rows)
    widths = [16, 10, 10, 14, 13, 14, 12, 14, 13, 14, 13, 14, 10, 12]
    for i, w in enumerate(widths[:len(headers)], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(headers)), n_rows + 1)


# ========== Read data ==========
print("=== 读取文件 ===")
df_july = pd.read_excel(BEFORE_INV, sheet_name='sheet1', dtype={'商品编码': str})
df_flow = pd.read_excel(NOW_FLOW, sheet_name='sheet1', dtype={'商品编码': str})
df_aug = pd.read_excel(AFTER_INV, sheet_name='sheet1', dtype={'商品编码': str})

print("7月账面库存: %d rows" % len(df_july))
print("8月出入库流水: %d rows" % len(df_flow))
print("8月账面库存: %d rows" % len(df_aug))

for df in [df_july, df_flow, df_aug]:
    for col in ['商品编码', '库存类型', '仓库', '货主']:
        if col in df.columns:
            df[col] = df[col].astype(str).str.strip()

# ========== Process flow data ==========
print("\n=== 处理8月出入库流水 ===")
df_flow['is_negative'] = df_flow['单据编号'].astype(str).apply(
    lambda x: any(x.startswith(p) for p in NEG_PREFIXES)
)
df_flow['出入库数量_符号'] = df_flow['出入库数量'].where(~df_flow['is_negative'], -df_flow['出入库数量'])
df_flow['实际出入库成本_符号'] = df_flow['实际出入库成本'].where(~df_flow['is_negative'], -df_flow['实际出入库成本'])
print("Negative rows: %d, Positive rows: %d" % (df_flow['is_negative'].sum(), (~df_flow['is_negative']).sum()))

flow_grouped = df_flow.groupby(GROUP_KEYS).agg(
    八月入库数=('出入库数量_符号', 'sum'),
    八月入库成本=('实际出入库成本_符号', 'sum')
).reset_index()
print("Flow grouped rows: %d" % len(flow_grouped))

july_grouped = df_july.groupby(GROUP_KEYS).agg(
    七月账面库存=('账面库存', 'sum'),
    七月实际成本=('实际入库成本合计', 'sum')
).reset_index()
print("July grouped rows: %d" % len(july_grouped))

aug_grouped = df_aug.groupby(GROUP_KEYS).agg(
    八月账面库存=('账面库存', 'sum'),
    八月实际成本=('实际入库成本合计', 'sum')
).reset_index()
print("August grouped rows: %d" % len(aug_grouped))

# ========== Merge all data ==========
print("\n=== 合并数据 ===")
merged = july_grouped.merge(flow_grouped, on=GROUP_KEYS, how='outer')
merged = merged.merge(aug_grouped, on=GROUP_KEYS, how='outer')
num_cols = ['七月账面库存', '七月实际成本', '八月入库数', '八月入库成本', '八月账面库存', '八月实际成本']
for col in num_cols:
    merged[col] = merged[col].fillna(0)

# ========== Forward calculation: 7月 + 流水 -> 8月 ==========
# 差异数 = (7月账面库存 + 8月入库数) - 8月账面库存
# 差异额 = (7月实际成本 + 8月入库成本) - 8月实际成本
merged['正向差异数'] = (merged['七月账面库存'] + merged['八月入库数'] - merged['八月账面库存']).round(2)
merged['正向差异额'] = (merged['七月实际成本'] + merged['八月入库成本'] - merged['八月实际成本']).round(2)
for col in ['七月实际成本', '八月入库成本', '八月实际成本']:
    merged[col] = merged[col].round(2)

# ========== Reverse calculation: 8月 - 流水 -> 7月 ==========
# 计算期初账面数 = 8月账面库存 - 8月入库数
# 计算期初成本 = 8月实际成本 - 8月入库成本
# 差异数 = 计算期初账面数 - 7月账面库存
# 差异额 = 计算期初成本 - 7月实际成本
merged['反向计算期初数'] = (merged['八月账面库存'] - merged['八月入库数']).round(2)
merged['反向计算期初成本'] = (merged['八月实际成本'] - merged['八月入库成本']).round(2)
merged['反向差异数'] = (merged['反向计算期初数'] - merged['七月账面库存']).round(2)
merged['反向差异额'] = (merged['反向计算期初成本'] - merged['七月实际成本']).round(2)

# ========== Sort ==========
merged['has_diff'] = (merged['正向差异数'].abs() >= 0.01) | (merged['正向差异额'].abs() >= 0.01)
merged = merged.sort_values(['has_diff', '商品编码'], ascending=[False, True]).reset_index(drop=True)

fwd_diff_count = merged['has_diff'].sum()
rev_has_diff = (merged['反向差异数'].abs() >= 0.01) | (merged['反向差异额'].abs() >= 0.01)
rev_diff_count = rev_has_diff.sum()

print("正向: Total %d rows, %d with diff" % (len(merged), fwd_diff_count))
print("反向: Total %d rows, %d with diff" % (len(merged), rev_diff_count))

# ========== Generate Excel ==========
print("\n=== 生成Excel文件 ===")
wb = Workbook()

# --- Sheet 1: 正向差异对比 (7月+流水->8月) ---
ws1 = wb.active
ws1.title = '正向差异对比(7月推8月)'
headers1 = ['商品编码', '库存类型', '仓库', '货主', '7月账面库存', '7月实际成本',
            '8月入库数', '8月入库成本', '8月账面库存', '8月实际成本', '差异数', '差异额']
data_rows1 = []
for _, row in merged.iterrows():
    data_rows1.append((
        row['商品编码'], row['库存类型'], row['仓库'], row['货主'],
        float(row['七月账面库存']), float(row['七月实际成本']),
        float(row['八月入库数']), float(row['八月入库成本']),
        float(row['八月账面库存']), float(row['八月实际成本']),
        float(row['正向差异数']), float(row['正向差异额']),
        bool(row['has_diff'])
    ))
write_sheet(ws1, headers1, data_rows1,
            int_cols=[5, 7, 9, 11], dec_cols=[6, 8, 10, 12], diff_col_start=11)

# --- Sheet 2: 反向差异对比 (8月-流水->7月) ---
ws2 = wb.create_sheet('反向差异对比(8月推7月)')
headers2 = ['商品编码', '库存类型', '仓库', '货主', '8月账面库存', '8月实际成本',
            '8月入库数', '8月入库成本', '计算期初账面数', '计算期初成本',
            '7月账面库存', '7月实际成本', '差异数', '差异额']
data_rows2 = []
for _, row in merged.iterrows():
    data_rows2.append((
        row['商品编码'], row['库存类型'], row['仓库'], row['货主'],
        float(row['八月账面库存']), float(row['八月实际成本']),
        float(row['八月入库数']), float(row['八月入库成本']),
        float(row['反向计算期初数']), float(row['反向计算期初成本']),
        float(row['七月账面库存']), float(row['七月实际成本']),
        float(row['反向差异数']), float(row['反向差异额']),
        bool(rev_has_diff[_])
    ))
write_sheet(ws2, headers2, data_rows2,
            int_cols=[5, 7, 9, 11, 13], dec_cols=[6, 8, 10, 12, 14], diff_col_start=13)

# --- Sheet 3: 汇总统计 ---
ws3 = wb.create_sheet('汇总统计')
summary_data = [
    ['【正向验证】7月+流水 → 推算8月，对比8月文件', '', ''],
    ['指标', '数值', ''],
    ['总账面数', int(len(merged)), 'int'],
    ['有差异账面数', int(fwd_diff_count), 'int'],
    ['无差异账面数', int(len(merged) - fwd_diff_count), 'int'],
    ['7月账面库存合计', float(merged['七月账面库存'].sum()), 'int'],
    ['7月实际成本合计', round(float(merged['七月实际成本'].sum()), 2), 'dec'],
    ['8月入库数合计', float(merged['八月入库数'].sum()), 'int'],
    ['8月入库成本合计', round(float(merged['八月入库成本'].sum()), 2), 'dec'],
    ['8月账面库存合计', float(merged['八月账面库存'].sum()), 'int'],
    ['8月实际成本合计', round(float(merged['八月实际成本'].sum()), 2), 'dec'],
    ['差异数合计', round(float(merged['正向差异数'].sum()), 2), 'dec'],
    ['差异额合计', round(float(merged['正向差异额'].sum()), 2), 'dec'],
    ['', '', ''],
    ['【反向验证】8月-流水 → 推算7月，对比7月文件', '', ''],
    ['指标', '数值', ''],
    ['总账面数', int(len(merged)), 'int'],
    ['有差异账面数', int(rev_diff_count), 'int'],
    ['无差异账面数', int(len(merged) - rev_diff_count), 'int'],
    ['8月账面库存合计', float(merged['八月账面库存'].sum()), 'int'],
    ['8月实际成本合计', round(float(merged['八月实际成本'].sum()), 2), 'dec'],
    ['8月入库数合计', float(merged['八月入库数'].sum()), 'int'],
    ['8月入库成本合计', round(float(merged['八月入库成本'].sum()), 2), 'dec'],
    ['计算期初账面数合计', float(merged['反向计算期初数'].sum()), 'int'],
    ['计算期初成本合计', round(float(merged['反向计算期初成本'].sum()), 2), 'dec'],
    ['7月账面库存合计', float(merged['七月账面库存'].sum()), 'int'],
    ['7月实际成本合计', round(float(merged['七月实际成本'].sum()), 2), 'dec'],
    ['差异数合计', round(float(merged['反向差异数'].sum()), 2), 'dec'],
    ['差异额合计', round(float(merged['反向差异额'].sum()), 2), 'dec'],
]

section_font = Font(name='Arial', bold=True, size=11, color='2B4570')
section_fill = PatternFill(fill_type='solid', fgColor='EAF2FF')

for row_idx, (label, value, fmt) in enumerate(summary_data, 1):
    cell_a = ws3.cell(row_idx, 1, label)
    cell_b = ws3.cell(row_idx, 2, value)
    if label.startswith('【'):
        cell_a.font = section_font
        cell_a.fill = section_fill
        cell_b.fill = section_fill
        ws3.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=2)
    elif label == '指标' or label == '':
        cell_a.font = header_font
        cell_b.font = header_font
        cell_a.fill = header_fill
        cell_b.fill = header_fill
        cell_a.alignment = header_align
        cell_b.alignment = header_align
    elif label:
        cell_a.font = data_font
        cell_b.font = data_font
        if fmt == 'int':
            cell_b.number_format = FMT_INT
        elif fmt == 'dec':
            cell_b.number_format = FMT_DEC
    cell_a.border = thin_border
    cell_b.border = thin_border

ws3.column_dimensions['A'].width = 24
ws3.column_dimensions['B'].width = 22

wb.save(OUTPUT)
print("输出文件: %s" % OUTPUT)
print("完成!")
