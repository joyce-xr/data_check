import os
import time
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

# ==================== 配置区 ====================
INPUT_FILE = r"D:\粤之星对账\26-8\奕宝\奕宝-8月-出入库流水.xlsx"
OUTPUT_DIR = r"D:\粤之星对账\26-8\奕宝"
# 分组维度：可增删调整，如 ["商品编码", "库存类型", "仓库", "库存组织"]
GROUP_COLUMNS = ["商品编码", "库存类型", "仓库"]
# 单据编号以下列前缀开头的，出入库数量和成本转为负数
NEGATIVE_PREFIXES = ("XSDD", "CGTH", "PKD")
# ================================================

OUTPUT_FILE = os.path.join(OUTPUT_DIR, f"出入库汇总_{time.strftime('%Y%m%d_%H%M%S')}.xlsx")


def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    print(f"读取文件: {INPUT_FILE}")
    print(f"分组维度: {' + '.join(GROUP_COLUMNS)}")
    wb = openpyxl.load_workbook(INPUT_FILE, read_only=True, data_only=True)
    ws = wb.active

    qty_summary = defaultdict(float)
    cost_summary = defaultdict(float)
    product_names = {}

    header = None
    col_map = {}
    row_count = 0

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            header = row
            for j, name in enumerate(header):
                if name:
                    col_map[name] = j
            required = ["单据编号", "商品名称", "出入库数量", "实际出入库成本"] + GROUP_COLUMNS
            for r in required:
                if r not in col_map:
                    print(f"错误: 缺少列 '{r}'")
                    wb.close()
                    return
            print(f"表头: {header}")
            continue

        doc_no = row[col_map["单据编号"]] or ""
        name = row[col_map["商品名称"]] or ""
        qty = row[col_map["出入库数量"]] or 0
        cost = row[col_map["实际出入库成本"]] or 0

        # 构建分组 key
        group_vals = []
        skip_row = False
        for gcol in GROUP_COLUMNS:
            val = row[col_map[gcol]]
            if val is None:
                val = ""
            group_vals.append(str(val).strip())
        if skip_row:
            continue

        # 第一个分组列（通常是商品编码）为空则跳过
        if not group_vals[0]:
            continue

        key = tuple(group_vals)

        # XSDD/CGTH/PKD 开头的数量和成本转负数
        doc_prefix = str(doc_no).strip().upper()
        if doc_prefix.startswith(NEGATIVE_PREFIXES):
            qty = -abs(qty)
            cost = -abs(cost)

        qty_summary[key] += qty
        cost_summary[key] += cost
        if name:
            product_names[group_vals[0]] = name

        row_count += 1
        if row_count % 5000 == 0:
            print(f"  已处理 {row_count} 行...")

    wb.close()
    print(f"共处理 {row_count} 行数据")
    print(f"分组数量: {len(qty_summary)} 个账面")

    # 写入结果 Excel
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "出入库汇总"

    header_style = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    cell_font = Font(name="微软雅黑", size=10)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    center_align = Alignment(horizontal="center", vertical="center")

    headers = ["序号"] + GROUP_COLUMNS + ["商品名称", "出入库总数", "实际出入库成本总额"]
    ws_out.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws_out.cell(row=1, column=col)
        cell.font = header_style
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    sorted_keys = sorted(qty_summary.keys())

    for idx, key in enumerate(sorted_keys, 1):
        total_qty = qty_summary[key]
        total_cost = cost_summary[key]
        name = product_names.get(key[0], "")
        ws_out.append([idx] + list(key) + [name, round(total_qty, 4), round(total_cost, 4)])
        for col in range(1, len(headers) + 1):
            cell = ws_out.cell(row=idx + 1, column=col)
            cell.font = cell_font
            cell.border = thin_border
            if col == 1 or col in range(2, 2 + len(GROUP_COLUMNS)):
                cell.alignment = center_align

    # 列宽
    col_widths = [6] + [18] * len(GROUP_COLUMNS) + [36, 14, 20]
    for i, w in enumerate(col_widths):
        ws_out.column_dimensions[chr(65 + i)].width = w

    wb_out.save(OUTPUT_FILE)
    print(f"\n汇总结果已保存: {OUTPUT_FILE}")
    print(f"共 {len(sorted_keys)} 个账面（{' + '.join(GROUP_COLUMNS)}）")

    # 打印预览
    print("\n--- 预览（前20条）---")
    preview_headers = ["序号"] + GROUP_COLUMNS + ["商品名称", "出入库总数", "成本总额"]
    preview_widths = [5] + [16] * len(GROUP_COLUMNS) + [24, 10, 12]
    fmt = "  ".join(f"{{{i}:<{w}}}" for i, w in enumerate(preview_widths))
    print(fmt.format(*preview_headers))
    for idx, key in enumerate(sorted_keys[:20], 1):
        total_qty = qty_summary[key]
        total_cost = cost_summary[key]
        name = product_names.get(key[0], "")[:22]
        vals = [idx] + list(key) + [name, round(total_qty, 4), round(total_cost, 2)]
        print(fmt.format(*[str(v) for v in vals]))


if __name__ == "__main__":
    main()
